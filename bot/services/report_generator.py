import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import os


STATUS_COLORS = {
    "Можно оплачивать": "C6EFCE",
    "Требуется согласование": "FFEB9C",
    "Материал не найден в базе": "FFC7CE",
    "Первое поступление": "BDD7EE",
}


def _apply_header_style(ws, row=1):
    for cell in ws[row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
        cell.alignment = Alignment(horizontal="center")


def _status_fill(status: str):
    color = STATUS_COLORS.get(status, "FFFFFF")
    return PatternFill("solid", fgColor=color)


def generate_excel_report(invoice_data: dict, check_results: list, savings: dict) -> bytes:
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # === Итог ===
    ws_summary = wb.create_sheet("Итог")
    ws_summary.append(["Параметр", "Значение"])
    _apply_header_style(ws_summary)
    ws_summary.append(["Номер счета", invoice_data.get("invoice_number", "")])
    ws_summary.append(["Дата счета", invoice_data.get("invoice_date", "")])
    ws_summary.append(["Поставщик", invoice_data.get("supplier", "")])
    ws_summary.append(["ИНН", invoice_data.get("inn", "")])
    ws_summary.append(["Сумма", invoice_data.get("total_amount", "")])
    ws_summary.append(["НДС", invoice_data.get("vat_amount", "")])
    ws_summary.append(["Валюта", invoice_data.get("currency", "RUB")])
    ws_summary.append([""])
    ws_summary.append(["Всего позиций", savings.get("total_items", 0)])
    ws_summary.append(["Можно оплачивать", savings.get("ok_count", 0)])
    ws_summary.append(["Требуют согласования", savings.get("overpriced_count", 0)])
    ws_summary.append(["Не найдены в базе", savings.get("not_found_count", 0)])
    ws_summary.append(["Потенциальная экономия (руб. без НДС)", savings.get("total_savings", 0)])
    ws_summary.column_dimensions["A"].width = 40
    ws_summary.column_dimensions["B"].width = 30

    # === Все позиции ===
    headers = ["Наименование", "Артикул", "Ед. изм.", "Кол-во",
               "Цена без НДС", "Цена с НДС", "Сумма", "Мин. цена в базе",
               "Отклонение %", "Экономия/ед.", "Экономия итого", "Статус",
               "Поставщик мин. цены"]
    ws_all = wb.create_sheet("Все позиции")
    ws_all.append(headers)
    _apply_header_style(ws_all)
    for item in check_results:
        row = [
            item.get("name", ""),
            item.get("article", ""),
            item.get("unit", ""),
            item.get("quantity", ""),
            item.get("price_no_vat", ""),
            item.get("price_with_vat", ""),
            item.get("amount", ""),
            item.get("min_price", ""),
            item.get("deviation_percent", ""),
            item.get("savings_per_unit", ""),
            item.get("savings_amount", ""),
            item.get("status", ""),
            item.get("min_supplier", ""),
        ]
        ws_all.append(row)
        status = item.get("status", "")
        for cell in ws_all[ws_all.max_row]:
            cell.fill = _status_fill(status)
    for col in ws_all.columns:
        ws_all.column_dimensions[col[0].column_letter].width = 18

    # === Завышенные позиции ===
    ws_over = wb.create_sheet("Завышенные позиции")
    ws_over.append(headers)
    _apply_header_style(ws_over)
    for item in check_results:
        if item.get("status") == "Требуется согласование":
            ws_over.append([
                item.get("name", ""), item.get("article", ""), item.get("unit", ""),
                item.get("quantity", ""), item.get("price_no_vat", ""),
                item.get("price_with_vat", ""), item.get("amount", ""),
                item.get("min_price", ""), item.get("deviation_percent", ""),
                item.get("savings_per_unit", ""), item.get("savings_amount", ""),
                item.get("status", ""), item.get("min_supplier", ""),
            ])
            for cell in ws_over[ws_over.max_row]:
                cell.fill = _status_fill("Требуется согласование")
    for col in ws_over.columns:
        ws_over.column_dimensions[col[0].column_letter].width = 18

    # === Позиции не найдены в базе ===
    ws_nf = wb.create_sheet("Позиции не найдены в базе")
    ws_nf.append(["Наименование", "Артикул", "Ед. изм.", "Кол-во", "Цена без НДС", "Сумма"])
    _apply_header_style(ws_nf)
    for item in check_results:
        if item.get("status") == "Материал не найден в базе":
            ws_nf.append([
                item.get("name", ""), item.get("article", ""), item.get("unit", ""),
                item.get("quantity", ""), item.get("price_no_vat", ""), item.get("amount", ""),
            ])
    for col in ws_nf.columns:
        ws_nf.column_dimensions[col[0].column_letter].width = 22

    # === Аналоги ===
    ws_analogs = wb.create_sheet("Аналоги")
    ws_analogs.append(["Наименование позиции", "Аналог", "Примечание"])
    _apply_header_style(ws_analogs)
    ws_analogs.append(["(данные об аналогах будут добавлены позже)", "", ""])

    # === Минимальные цены ===
    ws_min = wb.create_sheet("Минимальные цены")
    ws_min.append(["Наименование", "Мин. цена без НДС", "Поставщик мин. цены"])
    _apply_header_style(ws_min)
    seen = set()
    for item in check_results:
        name = item.get("name", "")
        if name not in seen and item.get("min_price") is not None:
            seen.add(name)
            supplier = item.get("min_supplier")
            ws_min.append([name, item.get("min_price", ""), str(supplier or "—")])
    for col in ws_min.columns:
        ws_min.column_dimensions[col[0].column_letter].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_pdf_report(invoice_data: dict, check_results: list, savings: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            rightMargin=2*cm, leftMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Heading1"],
                                 fontSize=16, alignment=1, spaceAfter=12)
    heading_style = ParagraphStyle("Heading", parent=styles["Heading2"],
                                   fontSize=12, spaceAfter=6)
    normal_style = styles["Normal"]
    normal_style.fontSize = 9

    story = []

    # Cover / Title
    story.append(Paragraph("ОТЧЕТ ПО ПРОВЕРКЕ СЧЕТА", title_style))
    story.append(Spacer(1, 0.5*cm))

    invoice_info = [
        ["Параметр", "Значение"],
        ["Номер счета", str(invoice_data.get("invoice_number") or "—")],
        ["Дата счета", str(invoice_data.get("invoice_date") or "—")],
        ["Поставщик", str(invoice_data.get("supplier") or "—")],
        ["ИНН", str(invoice_data.get("inn") or "—")],
        ["Сумма", str(invoice_data.get("total_amount") or "—")],
        ["НДС", str(invoice_data.get("vat_amount") or "—")],
        ["Валюта", str(invoice_data.get("currency") or "RUB")],
        ["Дата проверки", datetime.now().strftime("%d.%m.%Y %H:%M")],
    ]
    t = Table(invoice_info, colWidths=[6*cm, 10*cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BACKGROUND", (0, 1), (0, -1), colors.lightgrey),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.5*cm))

    # Summary
    story.append(Paragraph("Сводка", heading_style))
    summary_data = [
        ["Показатель", "Значение"],
        ["Всего позиций", str(savings.get("total_items", 0))],
        ["Можно оплачивать", str(savings.get("ok_count", 0))],
        ["Требуют согласования", str(savings.get("overpriced_count", 0))],
        ["Не найдены в базе", str(savings.get("not_found_count", 0))],
        ["Потенциальная экономия (руб. без НДС)", str(savings.get("total_savings", 0))],
    ]
    t2 = Table(summary_data, colWidths=[10*cm, 6*cm])
    t2.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    story.append(t2)
    story.append(Spacer(1, 0.5*cm))

    # All positions table
    story.append(Paragraph("Все позиции счета", heading_style))
    col_headers = ["Наименование", "Ед.", "Кол-во", "Цена", "Мин. цена", "Откл.%", "Статус"]
    table_data = [col_headers]
    for item in check_results:
        status = item.get("status", "")
        table_data.append([
            str(item.get("name", ""))[:40],
            str(item.get("unit", "") or ""),
            str(item.get("quantity", "") or ""),
            str(item.get("price_no_vat", "") or ""),
            str(item.get("min_price", "") or "—"),
            str(item.get("deviation_percent", "") or "—"),
            status[:25],
        ])

    t3 = Table(table_data, colWidths=[5*cm, 1.2*cm, 1.5*cm, 2*cm, 2*cm, 1.5*cm, 3.8*cm])

    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.Color(0.95, 0.95, 0.95)]),
    ]
    for i, item in enumerate(check_results, start=1):
        status = item.get("status", "")
        if status == "Требуется согласование":
            style_cmds.append(("BACKGROUND", (6, i), (6, i), colors.Color(1, 0.92, 0.6)))
        elif status == "Можно оплачивать":
            style_cmds.append(("BACKGROUND", (6, i), (6, i), colors.Color(0.78, 0.94, 0.8)))
        elif status == "Материал не найден в базе":
            style_cmds.append(("BACKGROUND", (6, i), (6, i), colors.Color(1, 0.78, 0.78)))
    t3.setStyle(TableStyle(style_cmds))
    story.append(t3)

    doc.build(story)
    return buf.getvalue()
