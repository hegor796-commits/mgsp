import os
from datetime import date, datetime
from typing import Optional
import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


class ExcelDatabase:
    def __init__(self, db_path: str):
        self.db_path = db_path

    def init_db(self):
        os.makedirs(os.path.dirname(self.db_path), exist_ok=True)
        if os.path.exists(self.db_path):
            return

        wb = Workbook()
        wb.remove(wb.active)

        # Материалы
        ws = wb.create_sheet("Материалы")
        ws.append(["ID", "Нормализованное наименование", "Единица измерения",
                   "Минимальная цена без НДС", "Поставщик минимальной цены",
                   "Дата последнего обновления", "Категория", "Активен"])

        # История цен
        ws2 = wb.create_sheet("История цен")
        ws2.append(["ID", "ID материала", "Цена без НДС", "Цена с НДС",
                    "Поставщик", "Номер счета", "Дата счета", "Дата добавления", "Пользователь"])

        # Аналоги
        ws3 = wb.create_sheet("Аналоги")
        ws3.append(["ID", "ID материала", "ID аналога", "Коэффициент замены", "Примечание"])

        # Поставщики
        ws4 = wb.create_sheet("Поставщики")
        ws4.append(["ID", "Наименование", "ИНН", "Контакт", "Email", "Телефон", "Рейтинг"])

        # Пользователи
        ws5 = wb.create_sheet("Пользователи")
        ws5.append(["Telegram ID", "ФИО", "Роль", "Email", "Активен", "Дата добавления"])
        ws5.append([0, "Администратор", "admin", "", "Да", date.today().strftime("%Y-%m-%d")])

        # Проверки счетов
        ws6 = wb.create_sheet("Проверки счетов")
        ws6.append(["ID", "Номер счета", "Дата счета", "Поставщик", "Сумма", "НДС",
                    "Валюта", "Дата проверки", "Пользователь", "Статус", "Потенциальная экономия"])

        # Журнал
        ws7 = wb.create_sheet("Журнал")
        ws7.append(["ID", "Дата", "Пользователь (Telegram ID)", "Действие", "Детали"])

        wb.save(self.db_path)

    def _load_wb(self):
        return load_workbook(self.db_path)

    def _save_wb(self, wb):
        wb.save(self.db_path)

    def get_user(self, telegram_id) -> Optional[dict]:
        try:
            wb = self._load_wb()
            ws = wb["Пользователи"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None and str(row[0]) == str(telegram_id):
                    return {
                        "telegram_id": row[0],
                        "name": row[1],
                        "role": row[2],
                        "email": row[3],
                        "active": row[4],
                        "date_added": row[5],
                    }
            return None
        except Exception:
            return None

    def get_material(self, normalized_name: str) -> Optional[dict]:
        try:
            wb = self._load_wb()
            ws = wb["Материалы"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[1] and row[1].lower() == normalized_name.lower():
                    return {
                        "id": row[0],
                        "name": row[1],
                        "unit": row[2],
                        "min_price": row[3],
                        "min_price_supplier": row[4],
                        "last_updated": row[5],
                        "category": row[6],
                        "active": row[7],
                    }
            return None
        except Exception:
            return None

    def search_materials(self, query: str) -> list:
        results = []
        try:
            wb = self._load_wb()
            ws = wb["Материалы"]
            query_lower = query.lower()
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[1] and query_lower in row[1].lower():
                    results.append({
                        "id": row[0],
                        "name": row[1],
                        "unit": row[2],
                        "min_price": row[3],
                        "min_price_supplier": row[4],
                        "last_updated": row[5],
                        "category": row[6],
                        "active": row[7],
                    })
        except Exception:
            pass
        return results

    def get_min_price(self, material_id) -> Optional[float]:
        try:
            wb = self._load_wb()
            ws = wb["Материалы"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if str(row[0]) == str(material_id):
                    return float(row[3]) if row[3] is not None else None
            return None
        except Exception:
            return None

    def get_analogs(self, material_id) -> list:
        results = []
        try:
            wb = self._load_wb()
            ws = wb["Аналоги"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if str(row[1]) == str(material_id):
                    results.append({
                        "id": row[0],
                        "material_id": row[1],
                        "analog_id": row[2],
                        "coefficient": row[3],
                        "note": row[4],
                    })
        except Exception:
            pass
        return results

    def add_material(self, data: dict) -> str:
        try:
            wb = self._load_wb()
            ws = wb["Материалы"]
            # Generate new ID
            max_id = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None:
                    try:
                        max_id = max(max_id, int(row[0]))
                    except (ValueError, TypeError):
                        pass
            new_id = max_id + 1
            ws.append([
                new_id,
                data.get("name", ""),
                data.get("unit", ""),
                data.get("min_price", None),
                data.get("min_price_supplier", ""),
                data.get("last_updated", date.today().strftime("%Y-%m-%d")),
                data.get("category", ""),
                data.get("active", "Да"),
            ])
            self._save_wb(wb)
            return str(new_id)
        except Exception as e:
            raise RuntimeError(f"Ошибка добавления материала: {e}")

    def update_min_price(self, material_id, price, supplier, invoice_num, invoice_date, user):
        try:
            wb = self._load_wb()
            ws = wb["Материалы"]
            for row in ws.iter_rows(min_row=2):
                if str(row[0].value) == str(material_id):
                    row[3].value = price
                    row[4].value = supplier
                    row[5].value = date.today().strftime("%Y-%m-%d")
                    break
            self._save_wb(wb)
            self.add_price_history({
                "material_id": material_id,
                "price_no_vat": price,
                "price_with_vat": None,
                "supplier": supplier,
                "invoice_num": invoice_num,
                "invoice_date": invoice_date,
                "user": user,
            })
        except Exception as e:
            raise RuntimeError(f"Ошибка обновления цены: {e}")

    def add_price_history(self, data: dict):
        try:
            wb = self._load_wb()
            ws = wb["История цен"]
            max_id = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None:
                    try:
                        max_id = max(max_id, int(row[0]))
                    except (ValueError, TypeError):
                        pass
            new_id = max_id + 1
            ws.append([
                new_id,
                data.get("material_id"),
                data.get("price_no_vat"),
                data.get("price_with_vat"),
                data.get("supplier", ""),
                data.get("invoice_num", ""),
                data.get("invoice_date", ""),
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                data.get("user", ""),
            ])
            self._save_wb(wb)
        except Exception as e:
            raise RuntimeError(f"Ошибка добавления истории цен: {e}")

    def add_invoice_check(self, data: dict):
        try:
            wb = self._load_wb()
            ws = wb["Проверки счетов"]
            max_id = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None:
                    try:
                        max_id = max(max_id, int(row[0]))
                    except (ValueError, TypeError):
                        pass
            new_id = max_id + 1
            ws.append([
                new_id,
                data.get("invoice_number", ""),
                data.get("invoice_date", ""),
                data.get("supplier", ""),
                data.get("total_amount"),
                data.get("vat_amount"),
                data.get("currency", "RUB"),
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                data.get("user", ""),
                data.get("status", ""),
                data.get("potential_savings"),
            ])
            self._save_wb(wb)
        except Exception as e:
            raise RuntimeError(f"Ошибка добавления проверки счета: {e}")

    def get_all_materials(self) -> list:
        results = []
        try:
            wb = self._load_wb()
            ws = wb["Материалы"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None:
                    results.append({
                        "id": row[0],
                        "name": row[1],
                        "unit": row[2],
                        "min_price": row[3],
                        "min_price_supplier": row[4],
                        "last_updated": row[5],
                        "category": row[6],
                        "active": row[7],
                    })
        except Exception:
            pass
        return results

    def log_action(self, telegram_id, action: str, details: str = ""):
        try:
            wb = self._load_wb()
            ws = wb["Журнал"]
            max_id = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None:
                    try:
                        max_id = max(max_id, int(row[0]))
                    except (ValueError, TypeError):
                        pass
            new_id = max_id + 1
            ws.append([
                new_id,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                telegram_id,
                action,
                details,
            ])
            self._save_wb(wb)
        except Exception:
            pass
