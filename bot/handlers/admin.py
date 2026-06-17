from aiogram import Router, F, Bot
from aiogram.types import Message, CallbackQuery, BufferedInputFile
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup

from bot.config import settings
from bot.keyboards import main_menu_keyboard, admin_keyboard, cancel_keyboard
from bot.services.ocr import detect_and_parse
from bot.services.ai_extractor import AIExtractor

import os
import tempfile

ai_extractor = AIExtractor(settings.OPENAI_API_KEY)

router = Router()


class AddUserStates(StatesGroup):
    waiting_for_telegram_id = State()
    waiting_for_name = State()
    waiting_for_role = State()


class ChangeRoleStates(StatesGroup):
    waiting_for_telegram_id = State()
    waiting_for_role = State()


class AddMaterialStates(StatesGroup):
    waiting_for_input = State()  # file or text describing material(s)


def is_admin(user: dict) -> bool:
    return user and user.get("Роль") == "admin"


@router.message(F.text == "⚙️ Администрирование")
async def handle_admin_menu(message: Message, user: dict):
    if not is_admin(user):
        await message.answer("У вас нет прав для выполнения этой операции.")
        return
    await message.answer(
        "⚙️ <b>Панель администратора</b>\n\nВыберите действие:",
        parse_mode="HTML",
        reply_markup=admin_keyboard()
    )


@router.callback_query(F.data == "admin_add_user")
async def handle_add_user_start(callback: CallbackQuery, state: FSMContext, user: dict):
    if not is_admin(user):
        await callback.answer("Нет прав доступа.", show_alert=True)
        return
    await callback.answer()
    await state.set_state(AddUserStates.waiting_for_telegram_id)
    await callback.message.answer(
        "Введите Telegram ID нового пользователя (числовой идентификатор):",
        reply_markup=cancel_keyboard()
    )


@router.message(F.text == "❌ Отмена")
async def handle_cancel_admin(message: Message, state: FSMContext, user: dict):
    await state.clear()
    role = user.get("Роль", "user") if user else "user"
    await message.answer("Действие отменено.", reply_markup=main_menu_keyboard(role))


@router.message(AddUserStates.waiting_for_telegram_id)
async def handle_add_user_id(message: Message, state: FSMContext):
    text = message.text.strip()
    if not text.isdigit():
        await message.answer("Введите корректный числовой Telegram ID.")
        return
    await state.update_data(new_telegram_id=int(text))
    await state.set_state(AddUserStates.waiting_for_name)
    await message.answer("Введите ФИО пользователя:")


@router.message(AddUserStates.waiting_for_name)
async def handle_add_user_name(message: Message, state: FSMContext):
    await state.update_data(new_name=message.text.strip())
    await state.set_state(AddUserStates.waiting_for_role)
    await message.answer(
        "Введите роль пользователя:\n"
        "• <code>user</code> — обычный пользователь\n"
        "• <code>снабженец</code> — снабженец\n"
        "• <code>admin</code> — администратор",
        parse_mode="HTML"
    )


@router.message(AddUserStates.waiting_for_role)
async def handle_add_user_role(message: Message, state: FSMContext, user: dict):
    role = message.text.strip().lower()
    valid_roles = ["user", "снабженец", "admin"]
    if role not in valid_roles:
        await message.answer(
            f"Некорректная роль. Допустимые значения: {', '.join(valid_roles)}"
        )
        return

    data = await state.get_data()
    new_id = data.get("new_telegram_id")
    new_name = data.get("new_name")
    await state.clear()

    try:
        from bot.main import db
        db.add_user(new_id, new_name, role)
        db.log_action(
            message.from_user.id,
            "Добавление пользователя",
            f"ID: {new_id}, Имя: {new_name}, Роль: {role}"
        )
        admin_role = user.get("Роль", "admin") if user else "admin"
        await message.answer(
            f"✅ Пользователь успешно добавлен:\n"
            f"Telegram ID: {new_id}\n"
            f"ФИО: {new_name}\n"
            f"Роль: {role}",
            reply_markup=main_menu_keyboard(admin_role)
        )
    except Exception as e:
        await message.answer(f"Ошибка при добавлении пользователя: {e}")


@router.callback_query(F.data == "admin_add_material")
async def handle_add_material_start(callback: CallbackQuery, state: FSMContext, user: dict):
    if not is_admin(user):
        await callback.answer("Нет прав доступа.", show_alert=True)
        return
    await callback.answer()
    await state.set_state(AddMaterialStates.waiting_for_input)
    await callback.message.answer(
        "Отправьте файл (PDF, Excel, Word, фото) или напишите список материалов текстом.\n"
        "ИИ сам извлечёт названия, единицы измерения и цены и добавит в базу.",
        reply_markup=cancel_keyboard()
    )


@router.callback_query(F.data == "admin_change_role")
async def handle_change_role_start(callback: CallbackQuery, state: FSMContext, user: dict):
    if not is_admin(user):
        await callback.answer("Нет прав доступа.", show_alert=True)
        return
    await callback.answer()
    await state.set_state(ChangeRoleStates.waiting_for_telegram_id)
    await callback.message.answer(
        "Введите Telegram ID пользователя, которому нужно изменить роль:",
        reply_markup=cancel_keyboard()
    )


@router.message(ChangeRoleStates.waiting_for_telegram_id)
async def handle_change_role_id(message: Message, state: FSMContext):
    text = message.text.strip()
    if not text.isdigit():
        await message.answer("Введите корректный числовой Telegram ID.")
        return

    from bot.main import db
    target = db.get_user(int(text))
    if target is None:
        await message.answer("Пользователь с таким Telegram ID не найден в базе.")
        return

    await state.update_data(target_id=int(text), target_name=target.get("ФИО", ""))
    await state.set_state(ChangeRoleStates.waiting_for_role)
    await message.answer(
        f"Пользователь: {target.get('ФИО', '—')}\n"
        f"Текущая роль: {target.get('Роль', '—')}\n\n"
        f"Введите новую роль:\n"
        "• <code>user</code> — обычный пользователь\n"
        "• <code>снабженец</code> — снабженец\n"
        "• <code>admin</code> — администратор",
        parse_mode="HTML"
    )


@router.message(ChangeRoleStates.waiting_for_role)
async def handle_change_role_set(message: Message, state: FSMContext, user: dict):
    role = message.text.strip().lower()
    valid_roles = ["user", "снабженец", "admin"]
    if role not in valid_roles:
        await message.answer(
            f"Некорректная роль. Допустимые значения: {', '.join(valid_roles)}"
        )
        return

    data = await state.get_data()
    target_id = data.get("target_id")
    target_name = data.get("target_name")
    await state.clear()

    try:
        from bot.main import db
        db.update_user_role(target_id, role)
        db.log_action(
            message.from_user.id,
            "Изменение роли пользователя",
            f"ID: {target_id}, Имя: {target_name}, Новая роль: {role}"
        )
        admin_role = user.get("Роль", "admin") if user else "admin"
        await message.answer(
            f"✅ Роль обновлена:\n"
            f"Пользователь: {target_name}\n"
            f"Telegram ID: {target_id}\n"
            f"Новая роль: {role}",
            reply_markup=main_menu_keyboard(admin_role)
        )
    except Exception as e:
        await message.answer(f"Ошибка при изменении роли: {e}")


@router.callback_query(F.data == "admin_list_users")
async def handle_list_users(callback: CallbackQuery, user: dict):
    if not is_admin(user):
        await callback.answer("Нет прав доступа.", show_alert=True)
        return
    await callback.answer()

    try:
        from bot.main import db
        from openpyxl import load_workbook
        wb = load_workbook(db.db_path)
        ws = wb["Пользователи"]
        headers = [cell.value for cell in ws[1]]
        users = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in row):
                continue
            users.append(dict(zip(headers, row)))

        if not users:
            await callback.message.answer("Список пользователей пуст.")
            return

        lines = ["👥 <b>Список пользователей:</b>\n"]
        for u in users:
            lines.append(
                f"• {u.get('ФИО', '—')} (ID: {u.get('Telegram ID', '—')})\n"
                f"  Роль: {u.get('Роль', '—')}, "
                f"Активен: {u.get('Активен', '—')}"
            )
        await callback.message.answer(
            "\n".join(lines),
            parse_mode="HTML"
        )
    except Exception as e:
        await callback.message.answer(f"Ошибка получения списка пользователей: {e}")


def can_dedup(user: dict) -> bool:
    role = user.get("Роль", "") if user else ""
    return role in ("admin", "снабженец", "Снабженец")


@router.message(F.text == "🧹 Очистить дубликаты в базе")
async def handle_dedup_text(message: Message, user: dict):
    if not can_dedup(user):
        await message.answer("У вас нет прав для выполнения этой операции.")
        return
    try:
        from bot.main import db
        removed = db.deduplicate_materials()
        await message.answer(f"Готово. Удалено дублирующихся позиций: {removed}")
    except Exception as e:
        await message.answer(f"Ошибка при очистке дубликатов: {e}")


@router.callback_query(F.data == "admin_reset_materials")
async def handle_reset_materials_confirm(callback: CallbackQuery, user: dict):
    if not is_admin(user):
        await callback.answer("Нет прав доступа.", show_alert=True)
        return
    await callback.answer()
    from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton
    await callback.message.answer(
        "⚠️ Это удалит ВСЕ материалы и историю цен из базы без возможности восстановления.\n"
        "Пользователи и журнал действий не затронуты.\n\n"
        "Подтвердите удаление:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text="✅ Да, удалить всё", callback_data="admin_reset_materials_confirm"),
            InlineKeyboardButton(text="❌ Отмена", callback_data="admin_reset_materials_cancel"),
        ]])
    )


@router.callback_query(F.data == "admin_reset_materials_confirm")
async def handle_reset_materials_do(callback: CallbackQuery, user: dict):
    if not is_admin(user):
        await callback.answer("Нет прав доступа.", show_alert=True)
        return
    await callback.answer()
    try:
        from bot.main import db
        removed = db.reset_materials()
        db.log_action(callback.from_user.id, "Полная очистка базы материалов", f"Удалено: {removed}")
        await callback.message.answer(f"✅ База материалов очищена. Удалено позиций: {removed}")
    except Exception as e:
        await callback.message.answer(f"Ошибка при очистке базы: {e}")


@router.callback_query(F.data == "admin_reset_materials_cancel")
async def handle_reset_materials_cancel(callback: CallbackQuery):
    await callback.answer()
    await callback.message.answer("Отменено.")


@router.callback_query(F.data == "admin_dedup")
async def handle_dedup(callback: CallbackQuery, user: dict):
    if not can_dedup(user):
        await callback.answer("Нет прав доступа.", show_alert=True)
        return
    await callback.answer()
    try:
        from bot.main import db
        removed = db.deduplicate_materials()
        await callback.message.answer(
            f"Готово. Удалено дублирующихся позиций: {removed}"
        )
    except Exception as e:
        await callback.message.answer(f"Ошибка при очистке дубликатов: {e}")


@router.message(F.text == "➕ Добавить данные в базу знаний")
async def handle_add_to_kb(message: Message, state: FSMContext, user: dict):
    if not is_admin(user) and user.get("Роль") not in ("снабженец", "Снабженец"):
        await message.answer("У вас нет прав для выполнения этой операции.")
        return
    await state.set_state(AddMaterialStates.waiting_for_input)
    await message.answer(
        "Отправьте файл (PDF, Excel, Word, фото) или напишите список материалов текстом.\n"
        "ИИ сам извлечёт названия, единицы измерения и цены и добавит в базу.",
        reply_markup=cancel_keyboard()
    )


@router.message(AddMaterialStates.waiting_for_input, F.document | F.photo)
async def handle_kb_file(message: Message, state: FSMContext, bot: Bot, user: dict):
    from bot.main import db
    await message.answer("Файл получен, обрабатываю...")
    try:
        if message.document:
            file_id = message.document.file_id
            file_name = message.document.file_name or "material_file"
        else:
            file_id = message.photo[-1].file_id
            file_name = "material_photo.jpg"

        os.makedirs(settings.STORAGE_PATH, exist_ok=True)
        file_path = os.path.join(settings.STORAGE_PATH, file_name)
        await bot.download(file_id, destination=file_path)

        text = detect_and_parse(file_path)
        if not text or len(text.strip()) < 5:
            await message.answer("Не удалось извлечь текст из файла. Попробуйте другой формат.")
            return

        await _add_materials_from_text(message, state, db, text, user)
    except Exception as e:
        await message.answer(f"Ошибка при обработке файла: {e}")
        await state.clear()


@router.message(AddMaterialStates.waiting_for_input, F.text & ~F.text.startswith("❌"))
async def handle_kb_text(message: Message, state: FSMContext, user: dict):
    from bot.main import db
    await message.answer("Анализирую текст...")
    await _add_materials_from_text(message, state, db, message.text, user)


async def _add_materials_from_text(message: Message, state: FSMContext, db, text: str, user: dict):
    await message.answer("ИИ извлекает материалы из текста...")
    invoice_data = ai_extractor.extract_invoice(text)
    extracted = invoice_data.get("items", [])
    supplier = invoice_data.get("supplier") or ""

    if not extracted:
        await message.answer("Не удалось найти материалы в тексте. Попробуйте другой формат.")
        await state.clear()
        role = user.get("Роль", "user") if user else "user"
        await message.answer("Главное меню:", reply_markup=main_menu_keyboard(role))
        return

    added = 0
    skipped = 0
    seen: set[str] = set()

    for mat in extracted:
        name = (mat.get("name") or "").strip()
        if not name:
            continue
        price = ai_extractor.resolve_price_no_vat(mat, invoice_data) or mat.get("price_with_vat")
        characteristics = {k: v for k, v in mat.items()
                           if k not in ("name", "unit", "quantity", "price_no_vat",
                                        "price_with_vat", "amount") and v}
        normalized = ai_extractor.normalize_name(name, characteristics)
        key = normalized.lower().strip()
        if key in seen:
            continue
        seen.add(key)

        existing = db.get_material(normalized)
        if existing is None:
            first_word = normalized.split()[0] if normalized.split() else normalized
            for candidate in db.search_materials(first_word)[:5]:
                cand_name = candidate.get("Нормализованное наименование", "")
                if ai_extractor.is_same_material(normalized, cand_name):
                    existing = candidate
                    break

        if existing is None:
            db.add_material({
                "Нормализованное наименование": normalized,
                "Единица измерения": mat.get("unit") or "шт",
                "Минимальная цена без НДС": price,
                "Поставщик минимальной цены": supplier,
            })
            added += 1
        else:
            skipped += 1

    db.log_action(message.from_user.id, "Добавление материалов через ИИ",
                  f"Добавлено: {added}, пропущено дублей: {skipped}")

    await state.clear()
    role = user.get("Роль", "user") if user else "user"
    await message.answer(
        f"Готово.\n\n✅ Добавлено новых материалов: {added}\n⏭ Пропущено (уже есть в базе): {skipped}",
        reply_markup=main_menu_keyboard(role)
    )
