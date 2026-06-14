from datetime import datetime, timedelta
from aiogram import Router, F
from aiogram.types import Message, CallbackQuery, BufferedInputFile
from aiogram.fsm.context import FSMContext

from bot.keyboards import main_menu_keyboard, period_keyboard

router = Router()


@router.message(F.text == "📊 Выгрузить отчет")
async def handle_export_report(message: Message, user: dict):
    await message.answer(
        "Выберите период для отчета:",
        reply_markup=period_keyboard()
    )


@router.callback_query(F.data.in_(["period_today", "period_week", "period_month"]))
async def handle_report_period(callback: CallbackQuery, user: dict):
    period = callback.data
    role = user.get("Роль", "user") if user else "user"

    if period == "period_today":
        label = "Сегодня"
        since = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    elif period == "period_week":
        label = "За последнюю неделю"
        since = datetime.now() - timedelta(days=7)
    else:
        label = "За последний месяц"
        since = datetime.now() - timedelta(days=30)

    await callback.answer(f"Формирую отчет: {label}")

    try:
        from bot.main import db
        telegram_id = callback.from_user.id

        # Get checks from DB
        from openpyxl import load_workbook
        wb = load_workbook(db.db_path)
        ws = wb["Проверки счетов"]
        headers = [cell.value for cell in ws[1]]

        user_checks = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in row):
                continue
            row_dict = dict(zip(headers, row))
            row_user = str(row_dict.get("Пользователь", ""))
            if row_user != str(telegram_id):
                continue
            # Check date
            check_date_str = row_dict.get("Дата проверки", "")
            try:
                check_date = datetime.strptime(str(check_date_str)[:19], "%Y-%m-%d %H:%M:%S")
                if check_date >= since:
                    user_checks.append(row_dict)
            except Exception:
                pass

        if not user_checks:
            await callback.message.answer(
                f"За период «{label}» проверок не найдено.",
                reply_markup=main_menu_keyboard(role)
            )
            return

        total_savings = sum(
            float(c.get("Потенциальная экономия") or 0) for c in user_checks
        )

        lines = [
            f"📊 <b>Отчет по проверкам счетов</b>",
            f"Период: {label}",
            f"Всего проверок: {len(user_checks)}",
            f"Суммарная потенциальная экономия: {total_savings:.2f} руб.",
            "",
        ]

        for i, check in enumerate(user_checks[-20:], 1):
            lines.append(
                f"{i}. Счет №{check.get('Номер счета', '—')} "
                f"от {check.get('Дата счета', '—')}\n"
                f"   Поставщик: {check.get('Поставщик', '—')}\n"
                f"   Экономия: {float(check.get('Потенциальная экономия') or 0):.2f} руб.\n"
                f"   Проверен: {str(check.get('Дата проверки', '—'))[:16]}"
            )

        await callback.message.answer(
            "\n".join(lines),
            parse_mode="HTML",
            reply_markup=main_menu_keyboard(role)
        )

    except Exception as e:
        await callback.message.answer(
            f"Ошибка формирования отчета: {e}",
            reply_markup=main_menu_keyboard(role)
        )
